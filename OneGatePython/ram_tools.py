"""RAM simulation and visualization helpers, kept separate from instruction emission."""

from __future__ import annotations

import colorsys
from collections.abc import Mapping, Sequence, Set
from html import escape
from pathlib import Path

Instruction = tuple[int, int, int]
Allocation = tuple[int, str | None, bool]


def simulate_ram(program: Sequence[Instruction], ram_size: int = 256) -> list[int]:
    """Run NAND instructions from zeroed RAM and return the resulting snapshot."""
    ram = [0] * ram_size
    for address, (a, b, dst) in enumerate(program):
        if any(not 0 <= cell < ram_size for cell in (a, b, dst)):
            raise ValueError(
                f"instruction {address} contains an address outside RAM: "
                f"a={a}, b={b}, dst={dst}"
            )
        ram[dst] = 1 - (ram[a] & ram[b])
    return ram


def render_ram_table(ram: Sequence[int], start: int = 0, end: int = 256) -> str | None:
    """Format simulated RAM as a markdown 16-column grid for [start, end)."""
    if not 0 <= start <= end <= len(ram):
        raise ValueError(f"RAM range must be within 0..{len(ram)}, got {start}..{end}")
    if start == end:
        return None

    first_row = start & ~0x0F
    last_row = (end - 1) & ~0x0F
    header = ["Addr", *[f"0x{offset:02X}" for offset in range(16)]]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * len(header)) + " |",
    ]
    for base in range(first_row, last_row + 1, 16):
        row = [f"0x{base:02X}", *(str(bit) for bit in ram[base:base + 16])]
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def write_ram_table(
    ram: Sequence[int], path: str | Path, start: int = 0, end: int = 256
) -> None:
    """Write the markdown RAM table if the requested range is non-empty."""
    table = render_ram_table(ram, start, end)
    if table is not None:
        with open(path, "w") as file:
            file.write(table + "\n")


def _span(base: int, size: int) -> str:
    return f"0x{base:02X}" if size == 1 else f"0x{base:02X}..0x{base + size - 1:02X}"


def _allocation_color(index: int) -> str:
    hue = (index * 0.618033988749895) % 1
    red, green, blue = colorsys.hsv_to_rgb(hue, 0.3, 0.98)
    return "".join(f"{round(channel * 245):02X}" for channel in (red, green, blue))




def write_memory_map_xlsx(
    path: str | Path,
    allocations: Mapping[int, Allocation],
    free_variables: Set[int],
    free_scratch: Set[int],
    protected_range: tuple[int, int],
    scratch_range: tuple[int, int],
    zero: int,
    one: int,
    ram_size: int = 256,
) -> None:
    """Write an Excel RAM map with colored, row-merged allocation labels."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import MergedCell
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ordered_allocations = sorted(allocations.items())
    allocation_colors = {
        base: _allocation_color(index)
        for index, (base, _) in enumerate(ordered_allocations)
    }
    allocation_cells: dict[int, tuple[int, str | None, int, bool]] = {}
    for base, (size, name, protected) in ordered_allocations:
        for bit_index in range(size):
            allocation_cells[base + bit_index] = (base, name, size, protected)

    fills = {
        "constant_zero": PatternFill("solid", fgColor="FFEEEEEE"),
        "constant_one": PatternFill("solid", fgColor="FFD9EAD3"),
        "unnamed": PatternFill("solid", fgColor="FFC9DAF8"),
        "free": PatternFill("solid", fgColor="FFFFFFFF"),
        "unused": PatternFill("solid", fgColor="FFF3F3F3"),
        "header": PatternFill("solid", fgColor="FFD9E1F2"),
    }
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("The workbook did not create its default worksheet")
    sheet.title = "RAM Allocation Map"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "B2"
    sheet.column_dimensions["A"].width = 9
    for column in range(2, 18):
        sheet.column_dimensions[get_column_letter(column)].width = 12

    headers = ["Addr", *[f"0x{offset:02X}" for offset in range(16)]]
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(1, column, value)
        cell.fill = fills["header"]
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    merged_ranges: list[tuple[int, int, int, PatternFill]] = []
    for base in range(0, ram_size, 16):
        row = base // 16 + 2
        address_cell = sheet.cell(row, 1, f"0x{base:02X}")
        address_cell.font = Font(bold=True)
        address = base
        row_end = min(base + 16, ram_size)
        while address < row_end:
            column = address - base + 2
            allocation = allocation_cells.get(address)
            if allocation is not None:
                allocation_base, name, size, protected = allocation
                segment_end = min(allocation_base + size, row_end)
                end_column = segment_end - base + 1
                kind = "variable" if protected else "scratch"
                fill = (
                    PatternFill("solid", fgColor=f"FF{allocation_colors[allocation_base]}")
                    if name is not None else fills["unnamed"]
                )
                label = name or ""
                cell = sheet.cell(row, column, label)
                if name is not None:
                    cell.comment = Comment(
                        f"{kind.capitalize()} allocation @ {_span(allocation_base, size)}",
                        "OneGate",
                    )
                cell.alignment = Alignment(horizontal="left", vertical="center")
                for fill_column in range(column, end_column + 1):
                    sheet.cell(row, fill_column).fill = fill
                if end_column > column:
                    merged_ranges.append((row, column, end_column, fill))
                address = segment_end
                continue

            cell = sheet.cell(row, column)
            if isinstance(cell, MergedCell):
                raise RuntimeError(f"RAM address 0x{address:02X} unexpectedly overlaps a merged cell")
            if address == zero:
                cell.value, cell.fill = "constant zero", fills["constant_zero"]
            elif address == one:
                cell.value, cell.fill = "constant one", fills["constant_one"]
            elif address in free_variables:
                cell.value, cell.fill = "", fills["free"]
            elif address in free_scratch:
                cell.value, cell.fill = "", fills["free"]
            elif protected_range[0] <= address <= protected_range[1]:
                cell.value, cell.fill = "protected", fills["unnamed"]
            elif scratch_range[0] <= address <= scratch_range[1]:
                cell.value, cell.fill = "scratch", fills["unnamed"]
            else:
                cell.value, cell.fill = "unused", fills["unused"]
            cell.alignment = Alignment(horizontal="center")
            address += 1

    for row, start_column, end_column, fill in merged_ranges:
        sheet.merge_cells(
            start_row=row,
            start_column=start_column,
            end_row=row,
            end_column=end_column,
        )
        for column in range(start_column, end_column + 1):
            sheet.cell(row, column).fill = fill

    workbook.save(path)
